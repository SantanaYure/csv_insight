# -*- coding: utf-8 -*-
"""Gera InsurMinds_Desafio4_Relatorio_Tecnico.docx a partir de report_content.py."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content import COVER, BLOCKS

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x17, 0x32, 0x4A)
BLUE = RGBColor(0x2E, 0x6C, 0x9E)
GRAY = RGBColor(0x5B, 0x67, 0x70)
LIGHT_BG = "EAF2F8"
HEADER_BG = "17324A"

ASSETS_DIR = os.environ["REPORT_ASSETS_DIR"]
OUT_PATH = os.environ["REPORT_OUT_DOCX"]


def set_cell_bg(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_col_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])
    for idx, col in enumerate(table.columns):
        col.width = Cm(widths_cm[idx])


def style_base(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for i, size in ((1, 17), (2, 13.5), (3, 11.5)):
        style = doc.styles[f"Heading {i}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(16 if i == 1 else 10)
        style.paragraph_format.space_after = Pt(8 if i == 1 else 5)
        if i == 1:
            style.paragraph_format.keep_with_next = True
            pPr = style.element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "4")
            bottom.set(qn("w:color"), "17324A")
            pBdr.append(bottom)
            pPr.append(pBdr)


def add_footer_page_numbers(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CSV Insight — InsurMinds — Desafio 4  |  página ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run2 = p.add_run()
    run2.font.size = Pt(8.5)
    run2.font.color.rgb = GRAY
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)


def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["program"])
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["subtitle"])
    run.font.size = Pt(13)
    run.font.color.rgb = BLUE
    p.paragraph_format.space_after = Pt(40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["challenge"])
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["title"])
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    p.paragraph_format.space_after = Pt(40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Projeto: {COVER['project_name']}")
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = NAVY

    for label, value in (
        ("Grupo", COVER["group_name"]),
        ("Integrantes", COVER["members"]),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}: {value}")
        run.font.size = Pt(11)
        run.font.color.rgb = GRAY if "PENDENTE" not in value else RGBColor(0xB0, 0x3A, 0x2E)
        run.font.italic = "PENDENTE" in value

    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Relatório técnico gerado em {COVER['date']}")
    run.font.size = Pt(10.5)
    run.font.color.rgb = GRAY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documento elaborado com base exclusiva no código do repositório do projeto.")
    run.font.size = Pt(9.5)
    run.font.italic = True
    run.font.color.rgb = GRAY

    doc.add_page_break()


def add_table(doc, spec):
    header = spec["header"]
    rows = spec["rows"]
    total_width_cm = 17.0
    fracs = spec.get("col_widths_frac") or [1.0 / len(header)] * len(header)
    widths_cm = [total_width_cm * f for f in fracs]

    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, HEADER_BG)

    for r_idx, row_values in enumerate(rows):
        for c_idx, value in enumerate(row_values):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            if r_idx % 2 == 1:
                set_cell_bg(cell, "F5F8FA")

    set_col_widths(table, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_bullets(doc, items, numbered=False):
    style_name = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style_name)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(item)
        run.font.size = Pt(10.3)


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "2E6C9E")
    pBdr.append(left)
    pPr.append(pBdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), LIGHT_BG)
    pPr.append(shd)
    run = p.add_run(text)
    run.font.italic = True
    run.font.size = Pt(9.7)
    run.font.color.rgb = RGBColor(0x33, 0x44, 0x55)


def add_image(doc, filename, caption):
    path = os.path.join(ASSETS_DIR, filename)
    doc.add_picture(path, width=Cm(14.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.font.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def build_body(doc):
    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h2":
            doc.add_heading(block[1], level=2)
        elif kind == "h3":
            doc.add_heading(block[1], level=3)
        elif kind == "p":
            p = doc.add_paragraph(block[1])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "bullets":
            add_bullets(doc, block[1], numbered=False)
        elif kind == "numbered":
            add_bullets(doc, block[1], numbered=True)
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "image":
            add_image(doc, block[1], block[2])
        elif kind == "note":
            add_note(doc, block[1])
        elif kind == "pagebreak":
            doc.add_page_break()


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style_base(doc)
    add_footer_page_numbers(doc)
    build_cover(doc)
    build_body(doc)

    doc.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    main()

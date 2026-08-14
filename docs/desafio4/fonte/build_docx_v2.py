# -*- coding: utf-8 -*-
"""Gera a versão em linguagem acessível do relatório (para importar no Google Docs)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content_v2 import COVER, BLOCKS

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x17, 0x32, 0x4A)
BLUE = RGBColor(0x2E, 0x6C, 0x9E)
GRAY = RGBColor(0x5B, 0x67, 0x70)
DARK_TEXT = RGBColor(0x22, 0x22, 0x22)
LIGHT_BG = "EAF2F8"
QA_BG = "F2F6F9"

ASSETS_DIR = os.environ["REPORT_ASSETS_DIR"]
OUT_PATH = os.environ["REPORT_OUT_DOCX"]


def shade(paragraph_or_cell, hex_color):
    el = paragraph_or_cell._p if hasattr(paragraph_or_cell, "_p") else paragraph_or_cell._tc
    pPr = el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def left_border(paragraph, color_hex, size=16):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)


def style_base(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK_TEXT
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.28

    for i, size in ((1, 17), (2, 13)):
        style = doc.styles[f"Heading {i}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(20 if i == 1 else 10)
        style.paragraph_format.space_after = Pt(10 if i == 1 else 6)
        if i == 1:
            style.paragraph_format.keep_with_next = True
            pPr = style.element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "4")
            bottom.set(qn("w:color"), "17324A")
            pBdr.append(bottom)
            pPr.append(pBdr)


def add_footer(doc):
    section = doc.sections[0]
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CSV Insight  ·  InsurMinds  ·  Desafio 4")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY


def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["program"])
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["subtitle"])
    run.font.size = Pt(13)
    run.font.color.rgb = BLUE
    p.paragraph_format.space_after = Pt(40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["challenge"])
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["title"])
    run.font.size = Pt(15)
    run.font.color.rgb = DARK_TEXT
    p.paragraph_format.space_after = Pt(40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Projeto: {COVER['project_name']}")
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = NAVY

    for label, value in (("Grupo", COVER["group_name"]), ("Integrantes", COVER["members"])):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}: {value}")
        run.font.size = Pt(11)
        run.font.color.rgb = GRAY

    for _ in range(7):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COVER["date"])
    run.font.size = Pt(10.5)
    run.font.color.rgb = GRAY

    doc.add_page_break()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(item)
        run.font.size = Pt(10.8)


def add_image(doc, filename, caption):
    path = os.path.join(ASSETS_DIR, filename)
    doc.add_picture(path, width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.font.italic = True
    run.font.size = Pt(9.2)
    run.font.color.rgb = GRAY
    p.paragraph_format.space_after = Pt(14)


def add_flow(doc, steps):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(16)
    shade(p, LIGHT_BG)
    text = "   →   ".join(steps)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = NAVY
    # padding via left/right indent illusion
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)


def add_qa(doc, question):
    container = doc.add_paragraph()
    container.paragraph_format.space_before = Pt(4)
    container.paragraph_format.space_after = Pt(2)
    shade(container, QA_BG)
    left_border(container, "2E6C9E", size=18)
    container.paragraph_format.left_indent = Cm(0.3)

    label = container.add_run("PERGUNTA")
    label.font.size = Pt(8.2)
    label.font.bold = True
    label.font.color.rgb = BLUE
    container.add_run("\n")
    qrun = container.add_run(question)
    qrun.font.size = Pt(11)
    qrun.font.bold = True
    qrun.font.color.rgb = NAVY

    container.add_run("\n\n")
    label2 = container.add_run("RESPOSTA")
    label2.font.size = Pt(8.2)
    label2.font.bold = True
    label2.font.color.rgb = BLUE
    container.add_run("\n")
    arun = container.add_run("(resposta a inserir)")
    arun.font.size = Pt(10.6)
    arun.font.italic = True
    arun.font.color.rgb = GRAY

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def build_body(doc):
    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h2":
            doc.add_heading(block[1], level=2)
        elif kind == "p":
            p = doc.add_paragraph(block[1])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "bullets":
            add_bullets(doc, block[1])
        elif kind == "image":
            add_image(doc, block[1], block[2])
        elif kind == "flow":
            add_flow(doc, block[1])
        elif kind == "qa":
            add_qa(doc, block[1])
        elif kind == "pagebreak":
            doc.add_page_break()


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    style_base(doc)
    add_footer(doc)
    build_cover(doc)
    build_body(doc)

    doc.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    main()
